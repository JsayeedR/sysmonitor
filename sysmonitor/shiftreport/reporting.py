from io import BytesIO
import html
import smtplib
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from monitor.models import NotificationGateway

from .models import ShiftImportantIssue
from .services import BDT, parse_email_list


XLSX_MIME = (
    'application/vnd.openxmlformats-officedocument.'
    'spreadsheetml.sheet'
)


def _display_name(user):
    if user.get_full_name().strip():
        return user.get_full_name().strip()
    return user.username


def _autosize(ws):
    for column in ws.columns:
        max_len = 0
        letter = get_column_letter(column[0].column)

        for cell in column:
            value = '' if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))

        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 55)


def build_shift_report_xlsx(report):
    """
    Build the Shift Report workbook completely in memory.

    Sheets:
      - Shift Summary
      - Outages
      - Generator Shifts
      - Important Issues
      - Previous Day (Night only)
    """
    wb = Workbook()

    ws = wb.active
    ws.title = 'Shift Summary'

    rows = [
        ['SysMonitor Shift Report', ''],
        ['Operational Date', report.report_date.isoformat()],
        ['Shift', report.get_shift_display()],
        [
            'Shift Window',
            (
                f'{report.shift_start.astimezone(BDT):%d-%m-%Y %H:%M}'
                f' to '
                f'{report.shift_end.astimezone(BDT):%d-%m-%Y %H:%M}'
            ),
        ],
        ['Prepared By', _display_name(report.prepared_by)],
        ['Handover To', report.handover_to.name],
        ['Handover Email', report.handover_to.email],
        ['', ''],
        ['PDB Outage Count', report.outage_summary.get('count', 0)],
        [
            'Total PDB Outage',
            report.outage_summary.get('total_duration', '0m'),
        ],
        [
            'Generator 01 Runtime',
            report.generator_summary.get('Gen-01', '0m'),
        ],
        [
            'Generator 02 Runtime',
            report.generator_summary.get('Gen-02', '0m'),
        ],
        [
            'Generator Mode Changes',
            report.generator_summary.get('mode_change_count', 0),
        ],
        [
            'Temperature Min',
            report.sensor_summary.get('temperature_min'),
        ],
        [
            'Temperature Max',
            report.sensor_summary.get('temperature_max'),
        ],
        [
            'Temperature Avg',
            report.sensor_summary.get('temperature_avg'),
        ],
        [
            'Humidity Min',
            report.sensor_summary.get('humidity_min'),
        ],
        [
            'Humidity Max',
            report.sensor_summary.get('humidity_max'),
        ],
        [
            'Humidity Avg',
            report.sensor_summary.get('humidity_avg'),
        ],
        ['', ''],
        ['Regular Activities', report.regular_activities],
        ['Issues Observed', report.issues_observed],
        ['Pending Handover', report.pending_handover],
        ['Important Notes', report.important_notes],
    ]

    if report.shift == 'NIGHT':
        rows.append(['MNOC Notes', report.mnoc_notes])

    for row in rows:
        ws.append(row)

    ws['A1'].font = Font(bold=True, size=14)

    for cell in ws['A']:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical='top')

    for cell in ws['B']:
        cell.alignment = Alignment(
            vertical='top',
            wrap_text=True,
        )

    _autosize(ws)

    # ----------------------------------------------------------
    # OUTAGES
    # ----------------------------------------------------------
    ws = wb.create_sheet('Outages')

    ws.append([
        'ID',
        'Start',
        'End',
        'Duration',
        'Generator',
        'Cycle Type',
        'Manual',
        'Alarm / Reason',
    ])

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in report.outage_summary.get('rows', []):
        ws.append([
            row.get('id'),
            row.get('start'),
            row.get('end'),
            row.get('duration'),
            row.get('generator'),
            row.get('cycle_type'),
            'Yes' if row.get('manual') else 'No',
            row.get('alarm_reason') or '',
        ])

    _autosize(ws)

    # ----------------------------------------------------------
    # GENERATOR MODE CHANGES
    # ----------------------------------------------------------
    ws = wb.create_sheet('Generator Shifts')

    ws.append([
        'ID',
        'Generator',
        'Switched At',
        'Note',
        'Added By',
    ])

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in report.generator_summary.get('mode_changes', []):
        ws.append([
            row.get('id'),
            row.get('generator'),
            row.get('switched_at'),
            row.get('note'),
            row.get('added_by'),
        ])

    _autosize(ws)

    # ----------------------------------------------------------
    # IMPORTANT ISSUES HISTORY
    # ----------------------------------------------------------
    ws = wb.create_sheet('Important Issues')

    ws.append([
        'Status',
        'Title',
        'Description',
        'Opened At',
        'Resolved At',
        'Added By',
    ])

    for cell in ws[1]:
        cell.font = Font(bold=True)

    issues = ShiftImportantIssue.objects.all().order_by(
        '-opened_at',
        '-id',
    )

    for issue in issues:
        ws.append([
            issue.get_status_display(),
            issue.title,
            issue.description,
            (
                issue.opened_at.astimezone(BDT).strftime(
                    '%d-%m-%Y %H:%M'
                )
                if issue.opened_at
                else ''
            ),
            (
                issue.resolved_at.astimezone(BDT).strftime(
                    '%d-%m-%Y %H:%M'
                )
                if issue.resolved_at
                else ''
            ),
            (
                _display_name(issue.added_by)
                if issue.added_by
                else ''
            ),
        ])

    _autosize(ws)

    # ----------------------------------------------------------
    # NIGHT PREVIOUS-DAY SUMMARY
    # ----------------------------------------------------------
    if report.shift == 'NIGHT':
        ws = wb.create_sheet('Previous Day')

        prev = report.previous_day_summary or {}
        outage = prev.get('outage', {})

        ws.append(['Previous Calendar Date', prev.get('date', '')])
        ws.append(['Outage Count', outage.get('count', 0)])
        ws.append([
            'Total Outage',
            outage.get('total_duration', '0m'),
        ])

        ws.append([])
        ws.append([
            'ID',
            'Start',
            'End',
            'Duration',
            'Generator',
            'Cycle Type',
        ])

        for cell in ws[5]:
            cell.font = Font(bold=True)

        for row in outage.get('rows', []):
            ws.append([
                row.get('id'),
                row.get('start'),
                row.get('end'),
                row.get('duration'),
                row.get('generator'),
                row.get('cycle_type'),
            ])

        _autosize(ws)

    stream = BytesIO()
    wb.save(stream)

    filename = (
        f'shift_report_{report.report_date.isoformat()}_'
        f'{report.shift.lower()}.xlsx'
    )

    return filename, stream.getvalue()


def build_recipient_lists(report, config):
    """
    Normal email:
      TO = handover officer + mandatory manager
      CC = mandatory CC + operator-added CC

    Duplicates are removed while preserving order.
    """
    to_list = parse_email_list(
        '\n'.join([
            report.handover_to.email or '',
            config.manager_email or '',
        ])
    )

    cc_list = parse_email_list(
        '\n'.join([
            config.mandatory_cc or '',
            report.additional_cc or '',
        ])
    )

    # If somebody added a TO address again in CC, keep it only in TO.
    to_keys = {x.casefold() for x in to_list}

    cc_list = [
        x
        for x in cc_list
        if x.casefold() not in to_keys
    ]

    return to_list, cc_list


def build_mnoc_recipient_lists(config):
    to_list = parse_email_list(config.mnoc_to)
    cc_list = parse_email_list(config.mnoc_cc)

    to_keys = {x.casefold() for x in to_list}

    cc_list = [
        x
        for x in cc_list
        if x.casefold() not in to_keys
    ]

    return to_list, cc_list


def build_subject(report):
    return (
        f'SysMonitor NOC Shift Report — '
        f'{report.get_shift_display()} — '
        f'{report.report_date:%d %b %Y}'
    )


def build_html_body(report):
    outage = report.outage_summary or {}
    generator = report.generator_summary or {}
    sensor = report.sensor_summary or {}

    prepared = html.escape(_display_name(report.prepared_by))
    handover = html.escape(report.handover_to.name)

    body = f"""
    <html>
      <body>
        <h2>SysMonitor NOC Shift Report</h2>

        <p>
          <strong>Date:</strong> {report.report_date:%d %b %Y}<br>
          <strong>Shift:</strong> {html.escape(report.get_shift_display())}<br>
          <strong>Prepared By:</strong> {prepared}<br>
          <strong>Handover To:</strong> {handover}
        </p>

        <h3>Automatic Operational Summary</h3>
        <ul>
          <li>PDB outages: {outage.get('count', 0)}</li>
          <li>Total PDB outage: {html.escape(str(outage.get('total_duration', '0m')))}</li>
          <li>Generator 01 runtime: {html.escape(str(generator.get('Gen-01', '0m')))}</li>
          <li>Generator 02 runtime: {html.escape(str(generator.get('Gen-02', '0m')))}</li>
          <li>Generator shifts: {generator.get('mode_change_count', 0)}</li>
          <li>Sensor readings: {sensor.get('reading_count', 0)}</li>
        </ul>

        <h3>Regular Activities</h3>
        <p>{html.escape(report.regular_activities or '—').replace(chr(10), '<br>')}</p>

        <h3>Issues Observed</h3>
        <p>{html.escape(report.issues_observed or '—').replace(chr(10), '<br>')}</p>

        <h3>Pending Handover</h3>
        <p>{html.escape(report.pending_handover or '—').replace(chr(10), '<br>')}</p>

        <h3>Important Notes</h3>
        <p>{html.escape(report.important_notes or '—').replace(chr(10), '<br>')}</p>

        <p>
          Detailed operational data and Important Issues history are attached
          in the Excel workbook.
        </p>
      </body>
    </html>
    """

    return body


def build_mnoc_html_body(report):
    previous = report.previous_day_summary or {}
    outage = previous.get('outage', {})

    return f"""
    <html>
      <body>
        <h2>SysMonitor MNOC Night Shift Report</h2>

        <p>
          <strong>Night Shift Date:</strong>
          {report.report_date:%d %b %Y}
        </p>

        <p>
          <strong>Previous Calendar Day:</strong>
          {html.escape(str(previous.get('date', '—')))}<br>
          <strong>Previous-day outage count:</strong>
          {outage.get('count', 0)}<br>
          <strong>Previous-day total outage:</strong>
          {html.escape(str(outage.get('total_duration', '0m')))}
        </p>

        <h3>MNOC Notes</h3>
        <p>
          {html.escape(report.mnoc_notes or '—').replace(chr(10), '<br>')}
        </p>
      </body>
    </html>
    """


def send_smtp_email(
    *,
    subject,
    html_body,
    to_list,
    cc_list=None,
    attachment=None,
):
    """
    Send through the existing SysMonitor Email NotificationGateway.

    attachment:
        (filename, bytes, mime_type)

    This function performs the real network send. Tests must mock SMTP.
    """
    cc_list = cc_list or []

    if not to_list:
        raise ValueError('No TO recipient configured.')

    try:
        gateway = NotificationGateway.objects.get(channel='email')
    except NotificationGateway.DoesNotExist:
        raise RuntimeError(
            'Email Notification Gateway is not configured.'
        )

    if not gateway.is_enabled:
        raise RuntimeError(
            'Email Notification Gateway is disabled.'
        )

    if not gateway.email_host:
        raise RuntimeError('Email SMTP host is not configured.')

    if not gateway.email_username:
        raise RuntimeError('Email SMTP username is not configured.')

    if not gateway.email_password:
        raise RuntimeError('Email SMTP password is not configured.')

    sender = (
        gateway.email_from.strip()
        or gateway.email_username.strip()
    )

    msg = MIMEMultipart()

    msg['Subject'] = subject
    msg['From'] = sender
    msg['To'] = ', '.join(to_list)

    if cc_list:
        msg['Cc'] = ', '.join(cc_list)

    msg.attach(
        MIMEText(
            html_body,
            'html',
            'utf-8',
        )
    )

    if attachment:
        filename, content, mime_type = attachment

        part = MIMEApplication(
            content,
            Name=filename,
        )

        part.add_header(
            'Content-Disposition',
            'attachment',
            filename=filename,
        )

        if mime_type:
            part.replace_header(
                'Content-Type',
                mime_type,
            )

        msg.attach(part)

    recipients = list(to_list) + list(cc_list)

    smtp = smtplib.SMTP(
        gateway.email_host,
        gateway.email_port,
        timeout=30,
    )

    try:
        smtp.ehlo()
        smtp.starttls()
        smtp.ehlo()

        smtp.login(
            gateway.email_username,
            gateway.email_password,
        )

        smtp.sendmail(
            sender,
            recipients,
            msg.as_string(),
        )

    finally:
        try:
            smtp.quit()
        except Exception:
            pass
