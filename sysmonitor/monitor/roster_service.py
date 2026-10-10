"""Duty-roster import and current/next shift resolution for SysMonitor.

The importer is intentionally format-tolerant for the monthly NOC roster that
uses a month/date header, engineer names down the left, and daily duty codes
(M/E/N/G/O/leave codes) across columns.
"""
from __future__ import annotations

import calendar
import hashlib
import re
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path

import pytz
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.datavalidation import DataValidation

from .models import DutyRoster, DutyRosterAssignment


BDT = pytz.timezone('Asia/Dhaka')
SHIFT_META = {
    'M': {'label': 'Morning', 'start': time(9, 0), 'end': time(17, 0)},
    'E': {'label': 'Evening', 'start': time(17, 0), 'end': time(23, 0)},
    'N': {'label': 'Night', 'start': time(23, 0), 'end': time(9, 0)},
}
KNOWN_CODES = {
    'M', 'E', 'N', 'G', 'O', 'CL', 'D', 'DL', 'EL', 'SL', 'ML', 'AL', 'L',
}


class RosterImportError(ValueError):
    pass


def _normalise_code(value):
    if value is None:
        return ''
    if isinstance(value, str):
        return re.sub(r'\s+', '', value).upper()
    return ''


def _month_from_cell(cell):
    value = cell.value
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)

    if isinstance(value, (int, float)) and 30000 <= value <= 80000:
        try:
            if is_date_format(cell.number_format):
                parsed = from_excel(value)
                return date(parsed.year, parsed.month, 1)
        except Exception:
            pass

    if not isinstance(value, str):
        return None

    text = value.strip()
    for fmt in (
        '%B %Y', '%b %Y', '%B-%Y', '%b-%Y', '%m/%Y', '%m-%Y',
        '%Y-%m', '%d/%m/%Y', '%d-%m-%Y',
    ):
        try:
            parsed = datetime.strptime(text, fmt)
            return date(parsed.year, parsed.month, 1)
        except ValueError:
            continue
    return None


def parse_roster_xlsx(file_bytes):
    """Parse an uploaded monthly roster and return normalized assignments.

    Returns a dict with month, assignments, warnings, engineer_count.
    Raises RosterImportError when the sheet cannot be safely understood.
    """
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)
    except Exception as exc:
        raise RosterImportError(f'Excel file could not be opened: {exc}') from exc

    ws = wb.active

    month = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12),
                            min_col=1, max_col=min(ws.max_column, 12)):
        for cell in row:
            month = _month_from_cell(cell)
            if month:
                break
        if month:
            break

    if month is None:
        raise RosterImportError(
            'Roster month/year could not be detected. Keep the month date '
            'near the top of the workbook (for example 01/10/2026).'
        )

    date_row = None
    label_col = None
    day_columns = {}

    for row_no in range(1, min(ws.max_row, 20) + 1):
        for col_no in range(1, min(ws.max_column, 12) + 1):
            value = ws.cell(row_no, col_no).value
            if not isinstance(value, str):
                continue
            label = re.sub(r'\s+', '', value).lower().rstrip('>')
            if label != 'date':
                continue

            candidate = {}
            for c in range(col_no + 1, ws.max_column + 1):
                raw = ws.cell(row_no, c).value
                try:
                    day = int(raw)
                except (TypeError, ValueError):
                    continue
                if 1 <= day <= 31:
                    candidate[c] = day

            if len(candidate) >= 20:
                date_row = row_no
                label_col = col_no
                day_columns = candidate
                break
        if date_row:
            break

    if date_row is None:
        raise RosterImportError(
            'The "Date >" row with day numbers 1-31 could not be detected.'
        )

    assignments = []
    warnings = []
    engineer_names = []
    last_day = calendar.monthrange(month.year, month.month)[1]

    # Engineer rows normally begin after the weekday row. Scan a bounded area
    # and accept only rows that look like real duty-code rows.
    for row_no in range(date_row + 2, min(ws.max_row, date_row + 20) + 1):
        name = ws.cell(row_no, label_col).value
        if not isinstance(name, str) or not name.strip():
            continue
        name = name.strip()

        raw_codes = []
        shift_code_count = 0
        known_count = 0
        for col_no, day in day_columns.items():
            code = _normalise_code(ws.cell(row_no, col_no).value)
            raw_codes.append((col_no, day, code))
            if code in KNOWN_CODES:
                known_count += 1
            if code in SHIFT_META:
                shift_code_count += 1

        if known_count < 5 or shift_code_count < 2:
            continue

        engineer_names.append(name)
        for _col_no, day, code in raw_codes:
            if not code or day > last_day:
                continue
            if len(code) > 12:
                warnings.append(
                    f'{name}, day {day}: ignored unrecognized duty value "{code[:24]}".'
                )
                continue
            if code not in KNOWN_CODES:
                warnings.append(
                    f'{name}, day {day}: imported non-standard code "{code}".'
                )
            assignments.append({
                'duty_date': date(month.year, month.month, day),
                'engineer_name': name,
                'duty_code': code,
                'source_row': row_no,
            })

    if not assignments or len(engineer_names) < 2:
        raise RosterImportError(
            'No engineer duty rows could be detected. The workbook must contain '
            'engineer names and daily duty codes such as M, E and N.'
        )

    # Operational shifts may not have multiple engineers assigned to the same
    # M/E/N slot. Missing slots are allowed but are surfaced as warnings so the
    # dashboard says Not assigned instead of guessing.
    slot_map = defaultdict(list)
    for item in assignments:
        if item['duty_code'] in SHIFT_META:
            slot_map[(item['duty_date'], item['duty_code'])].append(
                item['engineer_name']
            )

    errors = []
    for day in range(1, last_day + 1):
        duty_date = date(month.year, month.month, day)
        for code in ('M', 'E', 'N'):
            names = slot_map.get((duty_date, code), [])
            if len(names) > 1:
                errors.append(
                    f'{duty_date:%d/%m/%Y} {SHIFT_META[code]["label"]}: '
                    f'multiple officers assigned ({", ".join(names)}).'
                )
            elif not names:
                warnings.append(
                    f'{duty_date:%d/%m/%Y} {SHIFT_META[code]["label"]}: '
                    'no duty officer assigned.'
                )

    if errors:
        raise RosterImportError(' '.join(errors[:8]))

    return {
        'month': month,
        'assignments': assignments,
        'warnings': warnings,
        'engineer_count': len(set(engineer_names)),
    }


@transaction.atomic
def import_roster(file_bytes, filename, user=None):
    file_bytes = bytes(file_bytes)
    parsed = parse_roster_xlsx(file_bytes)
    month = parsed['month']

    source_filename = Path(filename or 'Roster.xlsx').name[:200]
    source_sha256 = hashlib.sha256(file_bytes).hexdigest()

    roster, created = DutyRoster.objects.get_or_create(
        month=month,
        defaults={
            'source_filename': source_filename,
            'imported_by': user,
            'import_warnings': parsed['warnings'],
            'original_file': file_bytes,
            'original_size': len(file_bytes),
            'original_sha256': source_sha256,
        },
    )

    if not created:
        roster.source_filename = source_filename
        roster.imported_by = user
        roster.import_warnings = parsed['warnings']
        roster.original_file = file_bytes
        roster.original_size = len(file_bytes)
        roster.original_sha256 = source_sha256
        roster.save()
        roster.assignments.all().delete()

    DutyRosterAssignment.objects.bulk_create([
        DutyRosterAssignment(
            roster=roster,
            duty_date=item['duty_date'],
            engineer_name=item['engineer_name'],
            duty_code=item['duty_code'],
            source_row=item['source_row'],
        )
        for item in parsed['assignments']
    ])

    return roster, created, parsed



def build_roster_template(month=None):
    """Return a server-compatible blank monthly NOC roster workbook."""
    if month is None:
        month = timezone.now().astimezone(BDT).date().replace(day=1)

    if isinstance(month, datetime):
        month = month.date()

    month = month.replace(day=1)
    last_day = calendar.monthrange(month.year, month.month)[1]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Duty Roster'

    ws['B3'] = 'NOC Duty Roster Schedule of Duty Officer'
    ws['B3'].font = Font(bold=True, size=14)

    ws['B4'] = datetime(month.year, month.month, 1)
    ws['B4'].number_format = 'mmmm yyyy'
    ws['B4'].font = Font(bold=True)

    ws['B6'] = 'Date >'
    ws['B6'].font = Font(bold=True)

    header_fill = PatternFill(
        fill_type='solid',
        fgColor='D9EAF7',
    )

    for day in range(1, last_day + 1):
        col = day + 2
        current = date(month.year, month.month, day)

        day_cell = ws.cell(6, col, day)
        day_cell.font = Font(bold=True)
        day_cell.fill = header_fill
        day_cell.alignment = Alignment(horizontal='center')

        weekday_cell = ws.cell(
            7,
            col,
            calendar.day_abbr[current.weekday()],
        )
        weekday_cell.alignment = Alignment(horizontal='center')

    ws['B7'] = 'Engineer Name'
    ws['B7'].font = Font(bold=True)

    # Five rows match the normal NOC roster layout. Replace these placeholders
    # with the actual engineer names before uploading.
    for offset in range(5):
        row = 8 + offset
        ws.cell(row, 2, f'Engineer Name {offset + 1}')

    # Duty-code dropdown makes the downloaded template difficult to mistype.
    allowed_codes = 'M,E,N,G,O,CL,D,DL,EL,SL,ML,AL,L'
    validation = DataValidation(
        type='list',
        formula1=f'"{allowed_codes}"',
        allow_blank=True,
    )
    validation.error = (
        'Use a supported duty code such as M, E, N, G, O, CL or D.'
    )
    validation.errorTitle = 'Invalid duty code'
    validation.prompt = 'Choose the duty code for this engineer/date.'
    validation.promptTitle = 'Duty code'
    ws.add_data_validation(validation)

    last_col = get_column_letter(last_day + 2)
    validation.add(f'C8:{last_col}12')

    ws['B15'] = 'Server Template Instructions'
    ws['B15'].font = Font(bold=True)

    instructions = [
        ('B16', 'Replace Engineer Name 1-5 with actual engineer names.'),
        ('B17', 'Enter one duty code for each engineer/date.'),
        ('B18', 'Do not delete the Date > row or day-number columns.'),
        ('B19', 'M = Morning 09:00-17:00'),
        ('B20', 'E = Evening 17:00-23:00'),
        ('B21', 'N = Night 23:00-09:00'),
        ('B22', 'G = General Duty'),
        ('B23', 'O = Off'),
        ('B24', 'CL/D/DL/EL/SL/ML/AL/L = leave-related codes'),
    ]
    for cell, value in instructions:
        ws[cell] = value

    ws.column_dimensions['B'].width = 34
    for col in range(3, last_day + 3):
        ws.column_dimensions[get_column_letter(col)].width = 5

    ws.freeze_panes = 'C8'

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()



def _ensure_bdt(now):
    if now is None:
        return timezone.now().astimezone(BDT)
    if timezone.is_naive(now):
        return BDT.localize(now)
    return now.astimezone(BDT)


def _slot_for(now):
    local = _ensure_bdt(now)
    today = local.date()
    current_time = local.time().replace(tzinfo=None)

    if current_time < time(9, 0):
        code = 'N'
        duty_date = today - timedelta(days=1)
        start_dt = BDT.localize(datetime.combine(duty_date, time(23, 0)))
        end_dt = BDT.localize(datetime.combine(today, time(9, 0)))
    elif current_time < time(17, 0):
        code = 'M'
        duty_date = today
        start_dt = BDT.localize(datetime.combine(today, time(9, 0)))
        end_dt = BDT.localize(datetime.combine(today, time(17, 0)))
    elif current_time < time(23, 0):
        code = 'E'
        duty_date = today
        start_dt = BDT.localize(datetime.combine(today, time(17, 0)))
        end_dt = BDT.localize(datetime.combine(today, time(23, 0)))
    else:
        code = 'N'
        duty_date = today
        start_dt = BDT.localize(datetime.combine(today, time(23, 0)))
        end_dt = BDT.localize(datetime.combine(today + timedelta(days=1), time(9, 0)))

    return {
        'code': code,
        'label': SHIFT_META[code]['label'],
        'duty_date': duty_date,
        'start_dt': start_dt,
        'end_dt': end_dt,
    }


def _next_slot(current):
    code = current['code']
    duty_date = current['duty_date']

    if code == 'M':
        next_code = 'E'
        next_date = duty_date
        start_dt = current['end_dt']
        end_dt = BDT.localize(datetime.combine(next_date, time(23, 0)))
    elif code == 'E':
        next_code = 'N'
        next_date = duty_date
        start_dt = current['end_dt']
        end_dt = BDT.localize(datetime.combine(next_date + timedelta(days=1), time(9, 0)))
    else:
        next_code = 'M'
        next_date = duty_date + timedelta(days=1)
        start_dt = current['end_dt']
        end_dt = BDT.localize(datetime.combine(next_date, time(17, 0)))

    return {
        'code': next_code,
        'label': SHIFT_META[next_code]['label'],
        'duty_date': next_date,
        'start_dt': start_dt,
        'end_dt': end_dt,
    }


def _assignment_for(slot):
    month = slot['duty_date'].replace(day=1)
    roster = DutyRoster.objects.filter(month=month).first()
    if roster is None:
        names = []
    else:
        names = list(
            roster.assignments
            .filter(duty_date=slot['duty_date'], duty_code=slot['code'])
            .order_by('engineer_name')
            .values_list('engineer_name', flat=True)
        )

    return roster, names


def _serialize_slot(slot):
    roster, names = _assignment_for(slot)
    start_dt = slot['start_dt']
    end_dt = slot['end_dt']

    return {
        'code': slot['code'],
        'label': slot['label'],
        'duty_date': slot['duty_date'].isoformat(),
        'engineers': names,
        'engineer': ' / '.join(names) if names else 'Not assigned',
        'assigned': bool(names),
        'roster_month': roster.month.strftime('%B %Y') if roster else None,
        'start_iso': start_dt.isoformat(),
        'end_iso': end_dt.isoformat(),
        'start_time': start_dt.strftime('%I:%M %p'),
        'end_time': end_dt.strftime('%I:%M %p'),
        'date_label': slot['duty_date'].strftime('%d %b %Y'),
    }


def resolve_current_and_next(now=None):
    current = _slot_for(now)
    next_slot = _next_slot(current)
    return {
        'generated_at': _ensure_bdt(now).isoformat(),
        'current': _serialize_slot(current),
        'next': _serialize_slot(next_slot),
    }
