from datetime import datetime
from io import BytesIO

import pytz
from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from .models import DutyRoster, DutyRosterAssignment, UserProfile
from .roster_service import import_roster, resolve_current_and_next


BDT = pytz.timezone('Asia/Dhaka')


def sample_roster_bytes():
    wb = Workbook()
    ws = wb.active
    ws['B3'] = 'NOC Duty Roster Schedule of Duty Officer'
    ws['B4'] = datetime(2026, 10, 1)
    ws['B6'] = 'Date >'
    for day in range(1, 32):
        ws.cell(6, day + 2, day)
        ws.cell(7, day + 2, 'Day')
    people = [
        ('Engineer Morning', 'M'),
        ('Engineer Evening', 'E'),
        ('Engineer Night', 'N'),
        ('Engineer General', 'G'),
    ]
    for row, (name, code) in enumerate(people, start=8):
        ws.cell(row, 2, name)
        for day in range(1, 32):
            ws.cell(row, day + 2, code)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


class DutyRosterTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser('rosteradmin', 'a@example.com', 'pass12345')
        self.viewer = User.objects.create_user(
            'viewer1',
            password='pass12345',
        )
        UserProfile.objects.create(
            user=self.viewer,
            role='viewer',
        )
        self.sample_bytes = sample_roster_bytes()
        import_roster(self.sample_bytes, 'Roster.xlsx', self.admin)

    def test_import_creates_assignments(self):
        self.assertEqual(DutyRosterAssignment.objects.filter(duty_code='M').count(), 31)
        self.assertEqual(DutyRosterAssignment.objects.filter(duty_code='E').count(), 31)
        self.assertEqual(DutyRosterAssignment.objects.filter(duty_code='N').count(), 31)

    def test_shift_boundaries(self):
        checks = [
            ((2026, 10, 10, 8, 59), 'N', 'Engineer Night'),
            ((2026, 10, 10, 9, 0), 'M', 'Engineer Morning'),
            ((2026, 10, 10, 16, 59), 'M', 'Engineer Morning'),
            ((2026, 10, 10, 17, 0), 'E', 'Engineer Evening'),
            ((2026, 10, 10, 22, 59), 'E', 'Engineer Evening'),
            ((2026, 10, 10, 23, 0), 'N', 'Engineer Night'),
        ]
        for values, code, engineer in checks:
            now = BDT.localize(datetime(*values))
            state = resolve_current_and_next(now)
            self.assertEqual(state['current']['code'], code)
            self.assertEqual(state['current']['engineer'], engineer)

    def test_night_rolls_to_next_morning(self):
        now = BDT.localize(datetime(2026, 10, 10, 23, 30))
        state = resolve_current_and_next(now)
        self.assertEqual(state['current']['code'], 'N')
        self.assertEqual(state['next']['code'], 'M')
        self.assertEqual(state['next']['duty_date'], '2026-10-11')

    def test_admin_page_is_admin_only(self):
        self.client.login(username='viewer1', password='pass12345')
        response = self.client.get('/duty-roster/')
        self.assertEqual(response.status_code, 403)
        self.client.logout()
        self.client.login(username='rosteradmin', password='pass12345')
        response = self.client.get('/duty-roster/')
        self.assertEqual(response.status_code, 200)

    def test_admin_can_download_original_and_viewer_cannot(self):
        roster = DutyRoster.objects.get(month='2026-10-01')

        self.assertEqual(roster.original_size, len(self.sample_bytes))
        self.assertEqual(bytes(roster.original_file), self.sample_bytes)
        self.assertEqual(len(roster.original_sha256), 64)

        self.client.login(
            username='rosteradmin',
            password='pass12345',
        )
        response = self.client.get(
            f'/duty-roster/{roster.id}/download/'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, self.sample_bytes)

        self.client.logout()
        self.client.login(
            username='viewer1',
            password='pass12345',
        )
        response = self.client.get(
            f'/duty-roster/{roster.id}/download/'
        )
        self.assertEqual(response.status_code, 403)

    def test_admin_can_download_server_template(self):
        self.client.login(
            username='rosteradmin',
            password='pass12345',
        )

        response = self.client.get(
            '/duty-roster/template.xlsx/?month=2026-11'
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument',
            response['Content-Type'],
        )

        wb = load_workbook(BytesIO(response.content))
        ws = wb.active

        self.assertEqual(
            ws['B3'].value,
            'NOC Duty Roster Schedule of Duty Officer',
        )
        self.assertEqual(ws['B6'].value, 'Date >')
        self.assertEqual(ws['C6'].value, 1)
        self.assertEqual(ws['AF6'].value, 30)
        self.assertIsNone(ws['AG6'].value)

        month_value = ws['B4'].value
        self.assertEqual(month_value.year, 2026)
        self.assertEqual(month_value.month, 11)

        self.client.logout()
        self.client.login(
            username='viewer1',
            password='pass12345',
        )
        response = self.client.get(
            '/duty-roster/template.xlsx/?month=2026-11'
        )
        self.assertEqual(response.status_code, 403)

    def test_viewer_can_read_current_api(self):
        self.client.login(username='viewer1', password='pass12345')
        response = self.client.get('/api/duty-roster/current/')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['ok'])
